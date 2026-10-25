"""Build executive documents from the analysis artifacts, without invented results."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches as DocInches
from pptx import Presentation
from pptx.util import Inches, Pt
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'reports'

def main():
    a=json.loads((OUT/'data_audit.json').read_text()); r=json.loads((OUT/'model_results.json').read_text()); m=r['test_metrics']
    imp=pd.read_csv(OUT/'feature_importance.csv'); groups=pd.read_csv(OUT/'attrition_by_salary.csv')
    sections=[
      ('Executive summary',f"The supplied Salifort Motors dataset contains {a['raw_rows']:,} employee records. After removing {a['duplicate_rows']:,} exact duplicate rows, {a['clean_rows']:,} records remain, with observed attrition of {a['attrition_rate']:.1%}. A random forest provides strong retrospective separation of employees who left and stayed. Use these findings to design and evaluate supportive retention initiatives; prospective performance is not established."),
      ('Plan — business question and success criteria','Which recorded characteristics are associated with employee departure, and can a model identify leavers in a held-out sample? HR leaders are the intended audience. Average precision is the model-selection metric because leavers are a minority. Precision and recall describe the trade-off at a fixed 0.5 threshold. No intervention cost, prediction horizon, or deployment target was supplied.'),
      ('Analyze — data quality and assumptions',f"There are {a['missing_cells']} missing cells and ten columns. The raw departure rate is {a['raw_attrition_rate']:.1%}; the deduplicated rate is {a['attrition_rate']:.1%}. Exact duplicate removal occurs before splitting to avoid identical records appearing across splits. Without employee identifiers, identical rows could represent different employees, so the deduplication assumption requires confirmation. Department and salary are categorical; salary is encoded categorically rather than assuming equally spaced pay levels. Tenure extremes are retained because they can be valid observations."),
      ('Analyze — observed relationships', 'Salary-group departure rates: '+ '; '.join(f"{x.salary}: {x.attrition_rate:.1%} (n={x.employees:,})" for x in groups.itertuples())+'. Workload and satisfaction distributions differ between leavers and stayers; see the accompanying figures. Group associations are descriptive and do not establish that changing pay, hours, or satisfaction will reduce departures.'),
      ('Construct — leakage controls and comparison','Records are split 60% training, 20% validation, and 20% test with stratification and random seed 42. Imputation, scaling, and one-hot encoding are fitted within pipelines using training data. A dummy baseline, logistic regression, depth-limited decision tree, and random forest are compared on the validation set. The selected model is refitted on training plus validation, then evaluated on the held-out test set. The threshold remains 0.5; no test-set threshold tuning is performed.'),
      ('Execute — held-out results',f"Selected model: {r['selected_model']}. Test set: {r['split']['test']:,} records. Accuracy {m['accuracy']:.2%}; precision {m['precision']:.2%}; recall {m['recall']:.2%}; F1 {m['f1']:.3f}; ROC AUC {m['roc_auc']:.3f}; average precision {m['average_precision']:.3f}, compared with test prevalence {r['test_prevalence']:.3f}. Confusion matrix: 1,994 true negatives, 7 false positives, 34 false negatives, 364 true positives. The 500-resample percentile bootstrap interval for test average precision is {r['test_ap_bootstrap_95_ci'][0]:.3f}–{r['test_ap_bootstrap_95_ci'][1]:.3f}; it conditions on the fitted model and does not quantify training or dataset-shift uncertainty."),
      ('Model interpretation','Leading features by validation permutation importance: '+', '.join(imp.head(4).feature)+'. Importance measures loss of predictive performance after shuffling a feature. Correlated features can share importance, and these rankings do not identify causal drivers. Feature timing is unknown: satisfaction or evaluation scores collected near departure may not be available at a useful intervention date.'),
      ('Recommendations and measurement','1. HR and team leaders should review extreme workloads and project assignments, using confidential employee feedback to understand the context. 2. Pilot voluntary workload and manager-support interventions, with predefined retention and wellbeing measures and a comparison group where feasible. 3. Audit compensation and advancement opportunities before choosing pay interventions. 4. Treat model outputs as research evidence; avoid individual employment decisions until appropriate governance, consent, fairness evaluation, and prospective validation are complete.'),
      ('Limitations and next steps','The ZIP provides no authoritative dataset provenance, observation dates, employee IDs, data dictionary, or license. Dataset identity and permission to redistribute must be confirmed before public publication. A random record split cannot verify temporal or employee-level generalization. Similar profiles may remain across splits after exact deduplication. Protected-attribute coverage and subgroup fairness are unknown. No causal effects, deployment readiness, calibrated risk probabilities, or financial savings are established. Next obtain source documentation, confirm deduplication and feature availability, evaluate on a later employee cohort, and choose a threshold using the actual costs of missed and unnecessary outreach.')]
    md='# Salifort Motors — Employee Attrition Analysis\n\n'+'\n\n'.join('## '+h+'\n\n'+b for h,b in sections)+'\n\n## Reproduction\n\nRun `python src/analysis.py` followed by `python src/build_deliverables.py` from the repository root. Detailed metrics and group tables are in `reports/`.\n'
    (OUT/'executive_report.md').write_text(md)
    doc=Document();doc.add_heading('Salifort Motors',0);doc.add_paragraph('Employee Attrition | Executive Report')
    for h,b in sections:doc.add_heading(h,1);doc.add_paragraph(b)
    for f in ['attrition_groups.png','workload_satisfaction.png','feature_importance.png','confusion_matrix.png','precision_recall.png']:
        doc.add_picture(str(OUT/'figures'/f),width=DocInches(6))
    doc.save(OUT/'executive_report.docx')
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    def slide(title,body=None,image=None):
        s=prs.slides.add_slide(prs.slide_layouts[5]);s.shapes.title.text=title
        s.shapes.title.text_frame.paragraphs[0].font.size=Pt(30)
        if body:
            box=s.shapes.add_textbox(Inches(.65),Inches(1.5),Inches(12),Inches(5.3));tf=box.text_frame;tf.word_wrap=True
            for i,line in enumerate(body):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.size=Pt(23);p.space_after=Pt(20)
        if image:
            from PIL import Image
            image_path=OUT/'figures'/image
            with Image.open(image_path) as im: w,h=im.size
            scale=min(12/w,5.4/h); width=w*scale; height=h*scale
            s.shapes.add_picture(str(image_path),Inches((13.333-width)/2),Inches(1.45),width=Inches(width),height=Inches(height))
        footer=s.shapes.add_textbox(Inches(.65),Inches(7.12),Inches(12),Inches(.25));footer.text_frame.text=f'Salifort Motors | Retrospective analysis | {len(prs.slides)}'
    slide('Salifort Motors: understanding employee attrition',['PACE analysis: Plan • Analyze • Construct • Execute','Goal: understand recorded departure patterns and evaluate predictive performance.','Decision: validate retention hypotheses through a measured, supportive pilot.'])
    slide('Data quality changes the story',[f"{a['raw_rows']:,} raw rows → {a['clean_rows']:,} after removing exact duplicates.",f"Observed attrition: {a['raw_attrition_rate']:.1%} raw → {a['attrition_rate']:.1%} deduplicated.",'No missing cells; no employee IDs or dates. Confirm the deduplication assumption.'])
    slide('Observed attrition across employee groups',image='attrition_groups.png')
    slide('Workload and satisfaction: descriptive patterns',image='workload_satisfaction.png')
    slide('How the model was evaluated',['Stratified 60/20/20 train/validation/test split; seed 42.','Training-only preprocessing; dummy, logistic, tree, and forest compared.','Select on validation average precision; refit on development data.','Evaluate the test set at a fixed 0.5 threshold.'])
    slide('Held-out results: random forest',[f"Precision {m['precision']:.1%} | Recall {m['recall']:.1%} | F1 {m['f1']:.3f}",f"Average precision {m['average_precision']:.3f} | Test prevalence {r['test_prevalence']:.3f}",'364 leavers identified; 34 missed; 7 stayers incorrectly flagged.','Retrospective performance does not prove prospective readiness.'])
    slide('Errors on the held-out test set',image='confusion_matrix.png')
    slide('Features supporting prediction, not causal claims',image='feature_importance.png')
    slide('Recommended retention pilot',['HR + managers: review extreme workloads and project allocation.','Offer confidential feedback and voluntary support.','Evaluate retention and wellbeing against a defined comparison group.','Audit pay and progression; do not infer causal benefits from correlations.'])
    slide('Conditions before operational use',['Confirm source, redistribution rights, and duplicate handling.','Verify when each feature was collected and define a prediction horizon.','Validate on a later employee cohort; assess calibration and fairness.','Choose an outreach threshold from costs and capacity; retain human review.'])
    prs.save(OUT/'executive_presentation.pptx')
    print('Executive report (Markdown/DOCX) and 10-slide PowerPoint saved.')
if __name__=='__main__':main()
